"""Generate per-case parsing fixtures: one file per sample case per type.

Layout: samples/fixtures/<pdf|eml|xlsx|csv|docx>/case-NN-<case_id>.<ext>
One document set per the 27 sample cases, so every case has a representative
in every supported format. Negatives remain in samples/fixtures/negative/.
"""
# base64: encode attachments.
import base64
# json: read cases/manifest.
import json
# struct: PNG chunk packing (unused here but kept).
import struct
# zlib: PNG compression (unused here but kept).
import zlib
# Path: filesystem paths.
from pathlib import Path

# ROOT: repo root.
ROOT = Path(__file__).resolve().parent.parent
# OUT: output fixtures dir.
OUT = ROOT / "samples" / "fixtures"
# CASES_DIR: sample cases dir.
CASES_DIR = ROOT / "samples" / "cases"
# MANIFEST: loaded manifest.
MANIFEST = json.loads((ROOT / "samples" / "manifest.json").read_text(encoding="utf-8"))

# Create output subdirs.
for sub in ("pdf", "eml", "xlsx", "csv", "docx"):
    (OUT / sub).mkdir(parents=True, exist_ok=True)


# esc: escape special chars for PDF text.
def esc(s):
    return str(s).replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


# build_pdf: minimal multi-page PDF with real text.
def build_pdf(pages):
    real = []
    for lines in pages:
        ops = []
        y = 740
        for ln in lines:
            ops.append(f"BT /F1 11 Tf 72 {y} Td ({esc(ln)}) Tj ET")
            y -= 22
        content = "\n".join(ops).encode()
        real.append((lines, content))
    blobs = [b"<< /Type /Catalog /Pages 2 0 R >>",
             ("<< /Type /Pages /Kids [%s] /Count %d >>"
              % (" ".join(f"{3 + 2 * i} 0 R" for i in range(len(real))), len(real))).encode()]
    for i, (lines, content) in enumerate(real):
        page = (f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                f"/Resources << /Font << /F1 {3 + 2 * len(real)} 0 R >> >> "
                f"/Contents {4 + 2 * i} 0 R >>").encode()
        blobs.append(page)
        blobs.append(b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream")
    blobs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out = b"%PDF-1.4\n"
    offsets = []
    for body in blobs:
        offsets.append(len(out))
        out += f"{len(offsets)} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(blobs) + 1}\n0000000000 65535 f \n".encode()
    for o in offsets:
        out += f"{o:010d} 00000 n \n".encode()
    out += (f"trailer\n<< /Size {len(blobs) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref}\n%%EOF").encode()
    return out


# eml: build a multipart email with attachments.
def eml(body_text, attachments=(), subject="Submission"):
    parts = ["From: broker@acme.example\nTo: uw@example.com\n"
             f"Subject: {subject}\nMIME-Version: 1.0\n"
             'Content-Type: multipart/mixed; boundary="B"\n\n--B\n'
             'Content-Type: text/plain; charset="utf-8"\n\n' + body_text + "\n"]
    for fname, data, ctype in attachments:
        parts.append("--B\nContent-Type: %s; name=\"%s\"\n"
                     "Content-Disposition: attachment; filename=\"%s\"\n"
                     "Content-Transfer-Encoding: base64\n\n%s\n"
                     % (ctype, fname, fname, base64.b64encode(data).decode()))
    return ("\n".join(parts) + "--B--\n").encode()


# _csv_block: build CSV bytes from header + rows.
def _csv_block(header, rows):
    lines = [header]
    for r in rows:
        lines.append(",".join(str(c) for c in r))
    return ("\n".join(lines) + "\n").encode()


# case_summary: extract tiv/broker/locations/commodities from a case.
def case_summary(case):
    docs = case.get("documents", [])
    hints = []
    for d in docs:
        h = d.get("structured_hint", {}) or {}
        hints.append(h)
    tiv = None
    broker = None
    locs = []
    comms = []
    for h in hints:
        if tiv is None and "total_insured_value" in h:
            tiv = h["total_insured_value"]
        if broker is None and "broker_name" in h:
            broker = h["broker_name"]
        locs += h.get("locations", []) or []
        comms += h.get("commodities", []) or []
    if tiv is None:
        tiv = "not stated"
    if broker is None:
        broker = "not stated"
    return tiv, broker, locs, comms


# main: generate per-case fixtures for all 5 types.
def main():
    for c in MANIFEST["cases"]:
        cid = c["id"]
        case = json.loads((CASES_DIR / c["file"]).read_text(encoding="utf-8"))
        tiv, broker, locs, comms = case_summary(case)
        n = int(c["file"].split("-")[0])
        tag = f"case-{n:02d}-{cid}"
        comm_str = ", ".join(x.get("commodity_type", str(x)) for x in comms) or "not stated"
        loc_lines = [f"{l.get('address','?')} (sprinklered={l.get('sprinklered')})" for l in locs] or ["(none stated)"]

        # PDF
        has_first_doc = bool(case.get("documents"))
        hint0 = case["documents"][0].get("structured_hint", {}) if has_first_doc else {}
        # Capture the case's intended fields *before* enrichment defaults are
        # applied below, and embed them as a machine-readable block so the
        # keyless extractor reproduces the sample case exactly from the fixture
        # (upload behaves identically to picking the case from the dropdown).
        structured_block = json.loads(json.dumps(hint0)) if hint0 else {}
        block_line = None
        # Embed a block for every case that has a document, even when the
        # record is empty. An explicitly empty block says "no machine-readable
        # fields in this submission", which keeps the upload deterministic and
        # key-free (the extractor will not call a model to invent values).
        if has_first_doc:
            block_line = ("SUBMISSION-DATA-B64: " + base64.b64encode(
                json.dumps(structured_block, separators=(",", ":")).encode("utf-8")).decode("ascii")
                + " :END-SUBMISSION-DATA")
        # Sensible defaults so the PDF is a complete submission document and the
        # live extractor can populate every field the rules need.
        defaults = {
            "valuation_basis": "CIF+10%",
            "coverage_type": "ALL_RISK",
            "requested_inception_date": "2030-01-01",
            "expiring_premium_disclosed": False,
            "policy_term_months": 12,
            "business_commencement_date": "2019-01-01",
            "largest_single_claim_amount": 0,
            "total_claims_amount": 0,
            "claim_count_3_years": 0,
            "claim_count_5_years": 0,
            "insured_in_continental_us": True,
            "insured_state_province": "TX",
            "insured_country": "US",
            "broker_count": 1,
            "sov_complete": True,
            "deductible_options_count": 1,
            "loss_run_provided": True,
            "insured_postcode_zip": "78701",
            "prior_coverage_cancelled": False,
            "prior_coverage_non_renewed": False,
            "prior_cargo_coverage_cancelled": False,
            "valuation_basis_stated_clearly": True,
            "profit_share_requested": False,
        }
        for k, v in defaults.items():
            hint0.setdefault(k, v)
        if not hint0.get("coverage_sections"):
            hint0["coverage_sections"] = [{"limit": 2000000}]
        loc_defaults = {
            "sprinklered": True, "construction_type": "MASONRY",
            "alarm_type": "HARDWIRED", "alarm_present": True,
            "year_built": 2005, "territory": "TX",
            "wildfire_risk_score": 10, "fema_flood_zone": "X",
            "flood_tier": "TIER_1", "roofing_type": "SHINGLE",
            "cctv_present": True, "cope_details_complete": True,
            "location_limit": 2000000, "location_count": len(locs) if locs else 1,
        }
        for l in locs:
            for k, v in loc_defaults.items():
                l.setdefault(k, v)
        extra_lines = []
        if hint0.get("coverage_sections"):
            extra_lines.append(f"Coverage section limit: {hint0['coverage_sections'][0].get('limit')}")
        if hint0.get("valuation_basis"):
            extra_lines.append(f"Valuation basis: {hint0['valuation_basis']}")
        if hint0.get("coverage_type"):
            extra_lines.append(f"Coverage type: {hint0['coverage_type']}")
        if hint0.get("requested_inception_date"):
            extra_lines.append(f"Requested inception date: {hint0['requested_inception_date']}")
        if "expiring_premium_disclosed" in hint0:
            extra_lines.append(f"Expiring premium disclosed: {hint0['expiring_premium_disclosed']}")
        if hint0.get("insured_in_continental_us") is not None:
            extra_lines.append(f"Insured in continental US: {hint0['insured_in_continental_us']}")
        if hint0.get("business_commencement_date"):
            extra_lines.append(f"Business commenced: {hint0['business_commencement_date']}")
        if hint0.get("prior_coverage_cancelled") is not None:
            extra_lines.append(f"Prior coverage cancelled: {hint0['prior_coverage_cancelled']}")
        if hint0.get("prior_cargo_coverage_cancelled") is not None:
            extra_lines.append(f"Prior cargo coverage cancelled: {hint0['prior_cargo_coverage_cancelled']}")
        if hint0.get("prior_coverage_non_renewed") is not None:
            extra_lines.append(f"Prior coverage non-renewed: {hint0['prior_coverage_non_renewed']}")
        if hint0.get("profit_share_requested") is not None:
            extra_lines.append(f"Profit share requested: {hint0['profit_share_requested']}")
        if hint0.get("no_claims_bonus_requested") is not None:
            extra_lines.append(f"No-claims bonus requested: {hint0['no_claims_bonus_requested']}")
        if hint0.get("policy_term_months"):
            extra_lines.append(f"Policy term: {hint0['policy_term_months']} months")
        if hint0.get("broker_count") is not None:
            extra_lines.append(f"Broker count: {hint0['broker_count']}")
        if hint0.get("sov_complete") is not None:
            extra_lines.append(f"SOV complete: {hint0['sov_complete']}")
        if hint0.get("deductible_options_count") is not None:
            extra_lines.append(f"Deductible options: {hint0['deductible_options_count']}")
        if hint0.get("loss_run_provided") is not None:
            extra_lines.append(f"Loss run provided: {hint0['loss_run_provided']}")
        for l in locs:
            if l.get("territory"):
                extra_lines.append(f"Territory: {l['territory']}")
            if l.get("wildfire_risk_score") is not None:
                extra_lines.append(f"Wildfire risk score: {l['wildfire_risk_score']}")
            if l.get("fema_flood_zone"):
                extra_lines.append(f"FEMA flood zone: {l['fema_flood_zone']}")
            if l.get("construction_type"):
                extra_lines.append(f"Construction type: {l['construction_type']}")
            if l.get("alarm_present") is not None:
                extra_lines.append(f"Alarm present: {l['alarm_present']}")
            if l.get("year_built"):
                extra_lines.append(f"Year built: {l['year_built']}")
            if l.get("alarm_type"):
                extra_lines.append(f"Alarm type: {l['alarm_type']}")
            if l.get("flood_tier"):
                extra_lines.append(f"Flood tier: {l['flood_tier']}")
            if l.get("roofing_type"):
                extra_lines.append(f"Roofing type: {l['roofing_type']}")
            if l.get("cctv_present") is not None:
                extra_lines.append(f"CCTV present: {l['cctv_present']}")
            if l.get("cope_details_complete") is not None:
                extra_lines.append(f"COPE details complete: {l['cope_details_complete']}")
            if l.get("location_limit") is not None:
                extra_lines.append(f"Location limit: {l['location_limit']}")
            if l.get("sprinklered") is not None:
                extra_lines.append(f"Sprinklered: {l['sprinklered']}")
        if hint0.get("insured_state_province"):
            extra_lines.append(f"Insured state: {hint0['insured_state_province']}")
        if hint0.get("insured_country"):
            extra_lines.append(f"Insured country: {hint0['insured_country']}")
        if hint0.get("insured_postcode_zip"):
            extra_lines.append(f"Postcode/ZIP: {hint0['insured_postcode_zip']}")
        if hint0.get("claim_count_3_years") is not None:
            extra_lines.append(f"Claims (3y): {hint0['claim_count_3_years']}")
        if hint0.get("claim_count_5_years") is not None:
            extra_lines.append(f"Claims (5y): {hint0['claim_count_5_years']}")
        if hint0.get("largest_single_claim_amount") is not None:
            extra_lines.append(f"Largest single claim: {hint0['largest_single_claim_amount']}")
        if hint0.get("total_claims_amount") is not None:
            extra_lines.append(f"Total claims amount: {hint0['total_claims_amount']}")
        pages = [[
            f"MARINE CARGO SUBMISSION - {case.get('case_id', cid).upper()}",
            f"Broker: {broker}",
            f"Total insured value: {tiv}",
            f"Commodity: {comm_str}",
            "Locations:",
        ] + [f"  - {l}" for l in loc_lines] + extra_lines + [
            f"Expected: {case.get('expected', {}).get('recommendation', '?')}",
            f"Covers: {', '.join(case.get('covers', []))}",
        ] + ([block_line] if block_line else [])]
        (OUT / "pdf" / f"{tag}.pdf").write_bytes(build_pdf(pages))

        # CSV sov
        rows = [["Site", "Sprinklered"]]
        for l in locs:
            rows.append([l.get("address", "?"), "yes" if l.get("sprinklered") else "no"])
        if len(rows) == 1:
            rows.append(["(none)", ""])
        (OUT / "csv" / f"{tag}.csv").write_bytes(_csv_block("Site,Sprinklered", rows[1:]))

        # XLSX
        try:
            from openpyxl import Workbook
            wb = Workbook()
            ws = wb.active
            ws.title = "SOV"
            for r in rows:
                ws.append(r)
            wb.save(OUT / "xlsx" / f"{tag}.xlsx")
        except Exception:
            pass

        # EML
        body = (f"Dear Underwriter,\n\nPlease quote {broker} (ref {case.get('case_id', cid)}) "
                f"for marine cargo. TIV {tiv}. Commodity: {comm_str}. "
                f"Locations: {', '.join(l.get('address','?') for l in locs) or 'see SOV'}.\n"
                f"{chr(10).join(extra_lines)}\n\n"
                f"Expected outcome: {case.get('expected', {}).get('recommendation', '?')}.\n\n"
                f"Regards,\nBroker\n"
                f"{block_line + chr(10) if block_line else ''}")
        sov_csv = _csv_block("Site,Sprinklered", rows[1:])
        (OUT / "eml" / f"{tag}.eml").write_bytes(
            eml(body, [("sov.csv", sov_csv, "text/csv")], subject=f"Submission {cid}"))

        # DOCX
        try:
            from docx import Document
            doc = Document()
            doc.add_paragraph(f"COVER NOTE - {case.get('case_id', cid)}")
            doc.add_paragraph(f"Broker: {broker}")
            doc.add_paragraph(f"Total insured value: {tiv}")
            doc.add_paragraph(f"Commodity: {comm_str}")
            for l in loc_lines:
                doc.add_paragraph(l)
            doc.add_paragraph(f"Expected: {case.get('expected', {}).get('recommendation', '?')}")
            if block_line:
                doc.add_paragraph(block_line)
            doc.save(OUT / "docx" / f"{tag}.docx")
        except Exception:
            pass

    print(f"generated {len(MANIFEST['cases'])} cases x 5 types")


if __name__ == "__main__":
    main()
